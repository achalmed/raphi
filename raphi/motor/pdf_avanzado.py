"""Lo avanzado (R7): tachar de verdad, resaltar, notas, estampar una firma o un sello, formularios, aplanar, comparar
dos versiones, verificar firmas digitales con su cadena, tablas a Excel y leer comprobantes (RUC, serie, fecha, total).

«Tachar» es redacción real: el texto de debajo se BORRA del archivo (no un rectángulo negro encima), y se comprueba
después que ya no se pueda extraer. Verificar firmas usa pyHanko: integridad (que nada cambió después de firmar),
firmante, fecha y, si se le dan raíces de confianza (`RAPHI_RAICES`: una carpeta con los certificados raíz, p. ej. los
de la ECERNEP del Perú), la cadena hasta ellas; sin raíces dice «cadena no verificada» en vez de inventarla.
"""
from __future__ import annotations

import datetime as dt
import difflib
import html
import os
import re
from pathlib import Path

from raphi.motor.base import ErrorOperacion, Parametro, Resultado, operacion
from raphi.motor.pdf_editar import _doc, _guardar
from raphi.motor.utiles import POSICIONES, color_rgb, interpretar_paginas, punto_en

PATRONES = {
    "dni": r"(?<!\d)\d{8}(?!\d)",
    "ruc": r"(?<!\d)(10|15|17|20)\d{9}(?!\d)",
    "correo": r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}",
    "celular": r"(?<![\d/-])9\d{2}[ -]?\d{3}[ -]?\d{3}(?![\d/-])",
}


def _buscar(pagina, terminos: list[str], patrones: list[str]) -> list:
    """Los rectángulos de cada aparición: los términos literales (sin distinguir mayúsculas) y los patrones."""
    import pymupdf
    cajas = []
    for t in terminos:
        cajas += pagina.search_for(t)
    if patrones:
        for palabra in pagina.get_text("words"):
            texto = palabra[4]
            if any(re.fullmatch(PATRONES[p], texto.strip(".,;:()")) for p in patrones):
                cajas.append(pymupdf.Rect(palabra[:4]))
    return cajas


def _lista(texto: str | None) -> list[str]:
    return [t.strip() for t in re.split(r"[\n;]", texto or "") if t.strip()]


@operacion("tachar", "Seguridad", "Tachar (redactar de verdad)",
           "Borra del archivo los términos o datos elegidos (DNI, RUC, correos, celulares) y los cubre; luego comprueba que ya no se puedan copiar.",
           sufijo="tachado", requiere=("py:pymupdf",),
           parametros=(Parametro("terminos", "Términos a tachar (uno por línea)", "multilinea", ""),
                       Parametro("dni", "Todos los DNI", "casilla", False), Parametro("ruc", "Todos los RUC", "casilla", False),
                       Parametro("correo", "Todos los correos", "casilla", False),
                       Parametro("celular", "Todos los celulares", "casilla", False),
                       Parametro("paginas", "Páginas (vacío = todas)", "paginas", ""),
                       Parametro("color", "Color del tachado", "color", "#000000")),
           palabras=("redactar", "censurar", "anonimizar", "ocultar datos", "privacidad"))
def tachar(ctx):
    terminos = _lista(ctx.p("terminos"))
    patrones = [p for p in PATRONES if ctx.p(p)]
    if not terminos and not patrones:
        raise ErrorOperacion("Diga qué tachar: términos o alguno de los datos (DNI, RUC, correo, celular).")
    doc = _doc(ctx.entrada)
    indices = interpretar_paginas(ctx.p("paginas"), doc.page_count)
    destino = ctx.salida()
    cuantos = 0
    for n in indices:
        pagina = doc[n]                     # una sola referencia: la anotación queda ligada a esta página
        cajas = _buscar(pagina, terminos, patrones)
        cuantos += len(cajas)
        if not ctx.simular:
            for c in cajas:
                pagina.add_redact_annot(c, fill=color_rgb(ctx.p("color")))
            pagina.apply_redactions()
    if ctx.simular:
        doc.close()
        return Resultado(plan=[f"{cuantos} apariciones a tachar → {destino.name}"])
    if not cuantos:
        doc.close()
        return Resultado(sin_cambio=True, mensajes=["No se encontró nada que tachar."])
    _guardar(doc, destino)
    restante = _doc(destino)
    quedan = sum(len(_buscar(restante[n], terminos, patrones)) for n in indices)
    restante.close()
    if quedan:
        destino.unlink()
        raise ErrorOperacion(f"Quedaron {quedan} apariciones legibles: no se entrega un tachado incompleto.")
    return Resultado([destino], {"tachados": cuantos}, [f"{cuantos} apariciones borradas del archivo"])


@operacion("resaltar", "Editar", "Resaltar términos", "Marca con resaltador cada aparición de los términos (anotaciones que se pueden quitar).",
           sufijo="resaltado", requiere=("py:pymupdf",),
           parametros=(Parametro("terminos", "Términos (uno por línea)", "multilinea", "", obligatorio=True),
                       Parametro("color", "Color", "color", "#FFE14D"), Parametro("paginas", "Páginas (vacío = todas)", "paginas", "")),
           palabras=("subrayar", "marcar", "highlight", "anotar"))
def resaltar(ctx):
    doc = _doc(ctx.entrada)
    indices = interpretar_paginas(ctx.p("paginas"), doc.page_count)
    destino = ctx.salida()
    total = 0
    for n in indices:
        pagina = doc[n]
        for c in _buscar(pagina, _lista(ctx.p("terminos")), []):
            total += 1
            if not ctx.simular:
                a = pagina.add_highlight_annot(c)
                a.set_colors(stroke=color_rgb(ctx.p("color")))
                a.update()
    if ctx.simular or not total:
        doc.close()
        return Resultado(plan=[f"{total} apariciones → {destino.name}"]) if ctx.simular else \
            Resultado(sin_cambio=True, mensajes=["No aparece ninguno de los términos."])
    return Resultado([_guardar(doc, destino)], {"resaltados": total})


@operacion("nota", "Editar", "Nota en una página", "Un cuadro de texto visible (o una nota adhesiva) en la página y posición elegidas.",
           sufijo="nota", requiere=("py:pymupdf",),
           parametros=(Parametro("texto", "Texto", "multilinea", "", obligatorio=True),
                       Parametro("pagina", "Página", "entero", 1, minimo=1),
                       Parametro("posicion", "Posición", "opcion", "arriba-derecha", opciones=POSICIONES),
                       Parametro("tipo", "Tipo", "opcion", "cuadro", opciones=(("cuadro", "Cuadro de texto visible"), ("adhesiva", "Nota adhesiva"))),
                       Parametro("autor", "Autor", "texto", "")),
           palabras=("comentario", "anotación", "observación", "post-it"))
def nota(ctx):
    import pymupdf
    doc = _doc(ctx.entrada)
    n = ctx.p("pagina") - 1
    if n >= doc.page_count:
        doc.close()
        raise ErrorOperacion(f"El documento tiene {doc.page_count} páginas.")
    destino = ctx.salida()
    if ctx.simular:
        doc.close()
        return Resultado(plan=[f"nota en la página {n + 1} → {destino.name}"])
    p = doc[n]
    lineas = ctx.p("texto").splitlines()
    ancho = min(260, max(pymupdf.get_text_length(l, fontsize=10) for l in lineas) + 16)
    alto = 14 * len(lineas) + 12
    x, y = punto_en(p.rect, ctx.p("posicion"), ancho, alto)
    caja = pymupdf.Rect(x, y, x + ancho, y + alto) * p.derotation_matrix
    if ctx.p("tipo") == "adhesiva":
        a = p.add_text_annot(caja.tl, ctx.p("texto"))
    else:
        a = p.add_freetext_annot(caja, ctx.p("texto"), fontsize=10, fill_color=(1, 0.98, 0.8), text_color=(0, 0, 0),
                                 rotate=p.rotation)
    a.set_info(title=ctx.p("autor") or "Raphi")
    a.update()
    return Resultado([_guardar(doc, destino)])


IMAGENES = (".png", ".jpg", ".jpeg", ".webp")


@operacion("estampar_imagen", "Editar", "Estampar firma o sello (imagen)",
           "Pone una imagen (la firma escaneada, un sello, un logo) en la página y posición elegidas. Es una firma visible, no una firma digital.",
           entrada="archivos", lote=False, extensiones=(".pdf",) + IMAGENES, sufijo="estampado", requiere=("py:pymupdf",),
           minimo_entradas=2,
           parametros=(Parametro("paginas", "Páginas", "paginas", "z", ayuda="z = la última; vacío = todas"),
                       Parametro("posicion", "Posición", "opcion", "abajo-derecha", opciones=POSICIONES),
                       Parametro("ancho", "Ancho (mm)", "decimal", 40, minimo=5, maximo=200),
                       Parametro("margen", "Margen (mm)", "decimal", 15, minimo=0, maximo=80)),
           palabras=("firma escaneada", "sello", "logo", "rúbrica", "imagen"))
def estampar_imagen(ctx):
    import pymupdf
    pdfs = [e for e in ctx.entradas if e.suffix.lower() == ".pdf"]
    imagenes = [e for e in ctx.entradas if e.suffix.lower() in IMAGENES]
    if len(imagenes) != 1 or not pdfs:
        raise ErrorOperacion("Elija uno o más PDF y UNA imagen (la firma o el sello).")
    from PIL import Image
    with Image.open(imagenes[0]) as img:
        proporcion = img.height / img.width
    ancho = ctx.p("ancho") / 25.4 * 72
    alto = ancho * proporcion
    salidas = []
    for pdf in pdfs:
        destino = ctx.salida(base=pdf)
        doc = _doc(pdf)
        indices = interpretar_paginas(ctx.p("paginas"), doc.page_count)
        if ctx.simular:
            doc.close()
            salidas.append(destino)
            continue
        for n in indices:
            p = doc[n]
            x, y = punto_en(p.rect, ctx.p("posicion"), ancho, alto, ctx.p("margen") / 25.4 * 72)
            p.insert_image(pymupdf.Rect(x, y, x + ancho, y + alto) * p.derotation_matrix, filename=str(imagenes[0]),
                           rotate=-p.rotation % 360 if p.rotation else 0, keep_proportion=True)
        salidas.append(_guardar(doc, destino))
    if ctx.simular:
        return Resultado(plan=[f"{imagenes[0].name} en {len(pdfs)} PDF → {', '.join(s.name for s in salidas)}"])
    return Resultado(salidas)


@operacion("formulario", "Editar", "Rellenar formulario",
           "Lista los campos de un formulario PDF y, si se dan valores, los rellena (y aplana, si se pide).",
           sufijo="relleno", requiere=("py:pymupdf",),
           parametros=(Parametro("valores", "Valores (JSON: {\"campo\": \"valor\"})", "multilinea", "",
                                 ayuda="vacío = solo listar los campos"),
                       Parametro("aplanar", "Aplanar (ya no se puede editar)", "casilla", False)),
           palabras=("formulario", "campos", "rellenar", "acroform", "llenar"))
def formulario(ctx):
    import json
    doc = _doc(ctx.entrada)
    campos = [{"nombre": w.field_name, "tipo": w.field_type_string, "valor": w.field_value, "pagina": p.number + 1}
              for p in doc for w in (p.widgets() or [])]
    if not campos:
        doc.close()
        return Resultado(sin_cambio=True, datos={"campos": []}, mensajes=["El PDF no tiene campos de formulario."])
    if not (ctx.p("valores") or "").strip():
        doc.close()
        return Resultado(datos={"campos": campos}, mensajes=[f"{c['nombre']} ({c['tipo']}) = {c['valor']!r}" for c in campos])
    try:
        valores = json.loads(ctx.p("valores"))
        assert isinstance(valores, dict)
    except (ValueError, AssertionError) as e:
        doc.close()
        raise ErrorOperacion('Los valores van en JSON: {"nombre": "Ana", "acepta": true}.') from e
    desconocidos = set(valores) - {c["nombre"] for c in campos}
    if desconocidos:
        doc.close()
        raise ErrorOperacion(f"El formulario no tiene: {', '.join(sorted(desconocidos))}.")
    destino = ctx.salida()
    if ctx.simular:
        doc.close()
        return Resultado(plan=[f"rellenar {len(valores)} campos → {destino.name}"])
    for p in doc:
        for w in p.widgets() or []:
            if w.field_name in valores:
                v = valores[w.field_name]
                w.field_value = (w.on_state() if v else "Off") if w.field_type_string == "CheckBox" else str(v)
                w.update()
    if ctx.p("aplanar"):
        doc.bake(annots=True, widgets=True)
    return Resultado([_guardar(doc, destino)], {"rellenados": len(valores)})


@operacion("aplanar", "Editar", "Aplanar anotaciones y formularios",
           "Convierte anotaciones, resaltados y campos en parte de la página: se ven igual en cualquier lector y ya no se editan.",
           sufijo="aplanado", requiere=("py:pymupdf",), palabras=("flatten", "fijar", "congelar"))
def aplanar(ctx):
    destino = ctx.salida()
    if ctx.simular:
        return Resultado(plan=[f"aplanar → {destino.name}"])
    doc = _doc(ctx.entrada)
    doc.bake(annots=True, widgets=True)
    return Resultado([_guardar(doc, destino)])


@operacion("comparar", "Información", "Comparar dos PDF",
           "Dice qué cambió entre dos versiones, página por página, con un informe HTML que marca lo quitado y lo añadido.",
           entrada="pdfs", lote=False, minimo_entradas=2, sufijo="comparacion", requiere=("py:pymupdf",),
           palabras=("diferencias", "versiones", "diff", "cambios", "proyecto de ley"))
def comparar(ctx):
    import pymupdf
    if len(ctx.entradas) != 2:
        raise ErrorOperacion("Elija exactamente dos PDF: la versión anterior y la nueva.")
    a, b = ctx.entradas
    destino = ctx.salida(nombre=f"comparacion_{a.stem}_{b.stem}", extension=".html")
    if ctx.simular:
        return Resultado(plan=[f"{a.name} ↔ {b.name} → {destino.name}"])

    def paginas(r):
        with pymupdf.open(r) as d:
            return [p.get_text() for p in d]
    pa, pb = paginas(a), paginas(b)
    secciones, distintas = [], []
    for i in range(max(len(pa), len(pb))):
        ta, tb = (pa[i] if i < len(pa) else ""), (pb[i] if i < len(pb) else "")
        if ta == tb:
            continue
        distintas.append(i + 1)
        wa, wb = ta.split(), tb.split()
        trozos = []
        for op, i1, i2, j1, j2 in difflib.SequenceMatcher(None, wa, wb, autojunk=False).get_opcodes():
            if op == "equal":
                trozos.append(html.escape(" ".join(wa[i1:i2])))
            if op in ("delete", "replace"):
                trozos.append(f"<del>{html.escape(' '.join(wa[i1:i2]))}</del>")
            if op in ("insert", "replace"):
                trozos.append(f"<ins>{html.escape(' '.join(wb[j1:j2]))}</ins>")
        secciones.append(f"<section><h2>Página {i + 1}</h2><p>{' '.join(trozos)}</p></section>")
    destino.write_text(
        '<!doctype html><html lang="es"><meta charset="utf-8"><title>Comparación</title><style>'
        "body{font:15px/1.6 system-ui,sans-serif;max-width:860px;margin:2em auto;padding:0 1em}"
        "del{background:#fdd;color:#900}ins{background:#dfd;color:#060;text-decoration:none}"
        "section{border-top:1px solid #ccc;margin-top:1.5em}</style>"
        f"<h1>{html.escape(a.name)} → {html.escape(b.name)}</h1>"
        f"<p>{len(pa)} → {len(pb)} páginas · {len(distintas)} con cambios · {dt.datetime.now():%d/%m/%Y %H:%M}</p>"
        + ("".join(secciones) or "<p>Las dos versiones dicen lo mismo.</p>") + "</html>", encoding="utf-8")
    return Resultado([destino], {"paginas_distintas": distintas, "paginas": [len(pa), len(pb)]},
                     [f"{len(distintas)} página(s) con cambios" if distintas else "Dicen lo mismo."])


# --- firmas digitales ----------------------------------------------------------------------------------------


def raices() -> list:
    """Los certificados raíz de confianza de `RAPHI_RAICES` (archivos .cer, .crt o .pem de una carpeta)."""
    carpeta = os.environ.get("RAPHI_RAICES")
    if not carpeta or not Path(carpeta).is_dir():
        return []
    from pyhanko.keys import load_certs_from_pemder
    certs = []
    for p in sorted(Path(carpeta).iterdir()):
        if p.suffix.lower() in (".cer", ".crt", ".pem", ".der"):
            certs += list(load_certs_from_pemder([str(p)]))
    return certs


def verificar(ruta: Path) -> list[dict]:
    """Cada firma: firmante, fecha, si el documento está íntegro desde que se firmó, si cubre todo el archivo y, con
    raíces, si la cadena llega a una de confianza."""
    from pyhanko.pdf_utils.reader import PdfFileReader
    from pyhanko.sign.validation import validate_pdf_signature
    from pyhanko_certvalidator import ValidationContext
    confianza = raices()
    contexto = ValidationContext(trust_roots=confianza, allow_fetching=False) if confianza else None
    salida = []
    with open(ruta, "rb") as f:
        lector = PdfFileReader(f, strict=False)
        for firma in lector.embedded_signatures:
            try:
                estado = validate_pdf_signature(firma, contexto) if contexto else validate_pdf_signature(firma)
                cert = firma.signer_cert
                salida.append({
                    "campo": firma.field_name, "firmante": cert.subject.human_friendly if cert else "",
                    "emisor": cert.issuer.human_friendly if cert else "",
                    "fecha": str(firma.self_reported_timestamp or ""), "integra": bool(estado.intact and estado.valid),
                    "cubre_todo": str(estado.coverage).endswith("ENTIRE_FILE"),
                    "cadena": ("de confianza" if estado.trusted else "no llega a una raíz de confianza") if confianza
                    else "no verificada (sin raíces: RAPHI_RAICES)",
                    "resumen": estado.summary() if hasattr(estado, "summary") else ""})
            except Exception as e:  # una firma ilegible no esconde a las demás
                salida.append({"campo": getattr(firma, "field_name", ""), "integra": False, "error": f"{type(e).__name__}: {e}"})
    return salida


@operacion("firmas", "Seguridad", "Verificar firmas digitales",
           "Quién firmó, cuándo, si el documento cambió después de firmarse y, con raíces de confianza (RAPHI_RAICES), si la cadena es válida.",
           requiere=("py:pyhanko",), palabras=("firma digital", "pades", "firma perú", "refirma", "certificado", "validar"))
def firmas(ctx):
    if ctx.simular:
        return Resultado(plan=[f"verificar las firmas de {ctx.entrada.name}"])
    lista = verificar(ctx.entrada)
    if not lista:
        return Resultado(datos={"firmas": []}, mensajes=["Sin firmas digitales."])
    mensajes = [f"{'✓' if f.get('integra') else '✗'} {f.get('firmante') or f.get('campo')} · {f.get('fecha', '')} · "
                f"{'íntegra' if f.get('integra') else 'NO íntegra'}{'' if f.get('cubre_todo', True) else ' (hay cambios después)'}"
                f" · cadena {f.get('cadena', '?')}" for f in lista]
    return Resultado(datos={"firmas": lista}, mensajes=mensajes)


# --- datos ----------------------------------------------------------------------------------------------------


@operacion("tablas_a_excel", "Convertir", "Tablas a Excel", "Encuentra las tablas del PDF y deja cada una en una hoja de Excel.",
           sufijo="tablas", requiere=("py:pymupdf", "py:openpyxl"),
           parametros=(Parametro("paginas", "Páginas (vacío = todas)", "paginas", ""),),
           palabras=("excel", "xlsx", "tabla", "datos", "planilla"))
def tablas_a_excel(ctx):
    import pymupdf
    from openpyxl import Workbook
    destino = ctx.salida(extension=".xlsx", nombre=f"{ctx.entrada.stem}_tablas")
    libro = Workbook()
    libro.remove(libro.active)
    with pymupdf.open(ctx.entrada) as doc:
        for n in interpretar_paginas(ctx.p("paginas"), doc.page_count):
            for i, t in enumerate(doc[n].find_tables().tables, 1):
                hoja = libro.create_sheet(f"p{n + 1}-t{i}"[:31])
                for fila in t.extract():
                    hoja.append([("" if c is None else c) for c in fila])
    if not libro.sheetnames:
        return Resultado(sin_cambio=True, mensajes=["No se encontraron tablas (si es un escaneo, pase antes por OCR)."])
    if ctx.simular:
        return Resultado(plan=[f"{len(libro.sheetnames)} tablas → {destino.name}"])
    libro.save(destino)
    return Resultado([destino], {"tablas": len(libro.sheetnames)})


SERIE = re.compile(r"\b([FBE]\w{3})\s*[-–]\s*0*(\d{1,8})\b")
FECHA = re.compile(r"\b(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})\b")
TOTAL = re.compile(r"(?:importe\s+)?total(?:\s+a\s+pagar)?\s*[:\s]*(?:S/\.?|PEN)?\s*([\d.,]+\d)", re.I)
RUC = re.compile(r"(?<!\d)((?:10|15|17|20)\d{9})(?!\d)")


def leer_datos(texto: str) -> dict:
    """RUC del emisor, serie y número, fecha y total de un comprobante peruano (boleta, factura) leído como texto."""
    def monto(t: str) -> float | None:
        t = t.replace(" ", "")
        if "," in t and "." in t:
            t = t.replace(",", "") if t.rfind(".") > t.rfind(",") else t.replace(".", "").replace(",", ".")
        elif "," in t:
            t = t.replace(",", ".")
        try:
            return float(t)
        except ValueError:
            return None
    ruc = RUC.search(texto)
    serie = SERIE.search(texto)
    fecha = FECHA.search(texto)
    totales = [m for m in (monto(x) for x in TOTAL.findall(texto)) if m is not None]
    iso = None
    if fecha:
        d, m, a = (int(x) for x in fecha.groups())
        a = a + 2000 if a < 100 else a
        try:
            iso = dt.date(a, m, d).isoformat()
        except ValueError:
            iso = None
    return {"ruc": ruc.group(1) if ruc else None, "serie": serie.group(1).upper() if serie else None,
            "numero": int(serie.group(2)) if serie else None, "fecha": iso, "total": max(totales) if totales else None}


@operacion("leer_comprobante", "Información", "Leer comprobante",
           "Saca RUC, serie y número, fecha y total de boletas y facturas (si es una foto o un escaneo, pase antes por OCR).",
           requiere=("py:pymupdf",), palabras=("boleta", "factura", "rendición", "viáticos", "ruc", "sunat"))
def leer_comprobante(ctx):
    import pymupdf
    if ctx.simular:
        return Resultado(plan=[f"leer {ctx.entrada.name}"])
    with pymupdf.open(ctx.entrada) as doc:
        texto = "\n".join(p.get_text() for p in doc)
    if len(texto.strip()) < 20:
        raise ErrorOperacion("El comprobante no tiene texto: páselo antes por «OCR».")
    d = leer_datos(texto)
    faltan = [k for k, v in d.items() if v is None]
    return Resultado(datos=d, mensajes=[" · ".join(f"{k}: {v}" for k, v in d.items() if v is not None)]
                     + ([f"No se encontró: {', '.join(faltan)}"] if faltan else []))
