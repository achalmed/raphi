"""Información: todo lo que se sabe de un PDF y el inventario de páginas (lo que hacía `page-counter`)."""
from __future__ import annotations

import datetime as dt
from pathlib import Path

import pikepdf

from raphi.motor.base import ErrorOperacion, Parametro, Resultado, operacion
from raphi.motor.utiles import tamano_legible


def informacion(ruta: Path) -> dict:
    datos: dict = {"archivo": ruta.name, "bytes": ruta.stat().st_size, "peso": tamano_legible(ruta.stat().st_size)}
    try:
        pdf = pikepdf.open(ruta)
    except pikepdf.PasswordError:
        return {**datos, "cifrado": True, "paginas": None, "mensaje": "Protegido con contraseña."}
    except pikepdf.PdfError as e:
        return {**datos, "valido": False, "mensaje": f"No abre como PDF: {e}"}
    with pdf:
        from raphi.motor.pdf_optimizar import es_pdfa
        from raphi.motor.pdf_seguridad import firmas
        tamanos: dict[str, int] = {}
        for p in pdf.pages:
            caja = [float(x) for x in p.mediabox]
            ancho, alto = (caja[2] - caja[0]) / 72 * 25.4, (caja[3] - caja[1]) / 72 * 25.4
            clave = _nombre_papel(ancho, alto)
            tamanos[clave] = tamanos.get(clave, 0) + 1
        info = {str(k).lstrip("/"): str(v) for k, v in pdf.docinfo.items()} if pdf.docinfo else {}
        datos.update({"valido": True, "paginas": len(pdf.pages), "version": pdf.pdf_version,
                      "cifrado": pdf.is_encrypted, "linearizado": pdf.is_linearized, "tamanos": tamanos,
                      "metadatos": info, "formulario": "/AcroForm" in pdf.Root and bool(pdf.Root.AcroForm.get("/Fields")),
                      "firmas": firmas(pdf), "pdfa": es_pdfa(ruta),
                      "marcadores": len(list(pdf.open_outline().root)) if "/Outlines" in pdf.Root else 0})
    try:
        from raphi.motor.pdf_contenido import texto_por_pagina
        largos = texto_por_pagina(ruta)
        datos["con_texto"] = sum(1 for n in largos if n >= 20)
    except Exception:
        pass
    return datos


PAPELES = {"A4": (210, 297), "A5": (148, 210), "A3": (297, 420), "Carta": (216, 279), "Oficio": (216, 330),
           "Legal": (216, 356)}


def _nombre_papel(ancho: float, alto: float) -> str:
    a, b = sorted((ancho, alto))
    for nombre, (x, y) in PAPELES.items():
        if abs(a - x) < 3 and abs(b - y) < 3:
            return nombre + (" apaisado" if ancho > alto else "")
    return f"{ancho:.0f}×{alto:.0f} mm"


@operacion("info", "Información", "Información del PDF",
           "Páginas, tamaño del papel, versión, cifrado, firmas, PDF/A, texto y metadatos.",
           palabras=("propiedades", "detalles", "páginas", "datos"))
def info(ctx):
    if ctx.simular:
        return Resultado(plan=[f"leer {ctx.entrada.name}"])
    d = informacion(ctx.entrada)
    if d.get("valido") is False:
        raise ErrorOperacion(d["mensaje"])
    mensajes = [f"{d.get('paginas')} páginas · {d['peso']} · PDF {d.get('version', '?')}"]
    if d.get("tamanos"):
        mensajes.append(" · ".join(f"{k}: {v}" for k, v in d["tamanos"].items()))
    return Resultado(datos=d, mensajes=mensajes)


# --- inventario de páginas (page-counter) ------------------------------------------------------------------------

ESTADO_OK = "OK"


def contar(ruta: Path) -> tuple[int, str]:
    try:
        with pikepdf.open(ruta) as pdf:
            return len(pdf.pages), ESTADO_OK
    except pikepdf.PasswordError:
        return 0, "Protegido"
    except Exception:
        return 0, "Error"


@operacion("inventario_paginas", "Información", "Inventario de páginas",
           "Cuenta las páginas de muchos PDF (archivos, carpetas o una fuente como «blogs del hub») y lo deja en un Excel con subtotales por grupo.",
           entrada="archivos", lote=False, extensiones=(), sufijo="inventario", requiere=("py:openpyxl",),
           parametros=(Parametro("fuente", "Fuente", "opcion", "archivos",
                                 opciones=(("archivos", "Los archivos y carpetas elegidos"),
                                           ("hub", "Blogs del hub (espacio de trabajo)"))),
                       Parametro("solo_index", "Solo los index.pdf", "casilla", False,
                                 ayuda="como page-counter: un PDF por publicación"),
                       Parametro("blogs", "Blogs (vacío = todos)", "texto", "", ayuda="solo con la fuente «hub»"),
                       Parametro("nombre", "Nombre del Excel", "texto", "")),
           minimo_entradas=0, palabras=("contar páginas", "page-counter", "excel", "conteo", "blogs"))
def inventario_paginas(ctx):
    from raphi.fuentes import agrupar
    grupos = agrupar(ctx.p("fuente"), ctx.entradas, solo_index=ctx.p("solo_index"),
                     blogs=[b.strip() for b in (ctx.p("blogs") or "").split(",") if b.strip()])
    archivos = sum(len(v) for _, v in grupos)
    if not archivos:
        raise ErrorOperacion("No se encontró ningún PDF en la fuente elegida.")
    nombre = ctx.p("nombre") or f"inventario-paginas_{dt.datetime.now():%Y-%m-%d_%H%M}"
    base = ctx.entradas[0] if ctx.entradas else Path.cwd() / "x"
    destino = ctx.salida(nombre=nombre, extension=".xlsx", base=base if base.is_file() else base / "x")
    if ctx.simular:
        return Resultado(plan=[f"{len(grupos)} grupos · {archivos} PDF → {destino.name}"])
    filas = []
    hechos = 0
    for grupo, rutas in grupos:
        resultados = []
        for raiz, ruta in rutas:
            ctx.avisar(hechos / archivos, ruta.name)
            n, estado = contar(ruta)
            resultados.append((str(ruta.relative_to(raiz)) if raiz else str(ruta), n, estado))
            hechos += 1
        filas.append((grupo, resultados))
    total = escribir_excel(destino, filas, ctx.p("solo_index"))
    return Resultado([destino], {"grupos": len(filas), "archivos": archivos, "paginas": total},
                     [f"{archivos} PDF · {total} páginas en {len(filas)} grupos"])


def escribir_excel(destino: Path, filas: list[tuple[str, list[tuple[str, int, str]]]], solo_index: bool) -> int:
    """El mismo libro que page-counter: hoja «Conteo de Páginas» con un bloque por grupo y su subtotal, el total
    general y la hoja «Información»."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    cab_f, cab_r = Font(bold=True, size=12, color="FFFFFF"), PatternFill("solid", start_color="2E5090", end_color="2E5090")
    gru_f, gru_r = Font(bold=True, size=11, color="FFFFFF"), PatternFill("solid", start_color="4472C4", end_color="4472C4")
    tot_f, tot_r = Font(bold=True, size=11), PatternFill("solid", start_color="E7E6E6", end_color="E7E6E6")
    borde = Border(*(Side(style="thin"),) * 4)
    centro = Alignment(horizontal="center", vertical="center")
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Conteo de Páginas"
    for celda, texto in {"A1": "Grupo", "B1": "Ruta del Archivo", "C1": "Número de Páginas", "D1": "Estado"}.items():
        hoja[celda] = texto
        hoja[celda].font, hoja[celda].fill, hoja[celda].alignment, hoja[celda].border = cab_f, cab_r, centro, borde
    fila, total, archivos = 3, 0, 0
    for grupo, resultados in filas:
        hoja[f"A{fila}"] = grupo.upper()
        hoja[f"A{fila}"].font, hoja[f"A{fila}"].fill = gru_f, gru_r
        hoja.merge_cells(f"A{fila}:D{fila}")
        fila += 1
        sub = 0
        for ruta, n, estado in resultados:
            hoja[f"B{fila}"], hoja[f"C{fila}"], hoja[f"D{fila}"] = ruta, n if estado == ESTADO_OK else 0, estado
            hoja[f"C{fila}"].alignment = hoja[f"D{fila}"].alignment = Alignment(horizontal="center")
            sub += n if estado == ESTADO_OK else 0
            fila += 1
        hoja[f"B{fila}"], hoja[f"C{fila}"], hoja[f"D{fila}"] = f"SUBTOTAL {grupo}", sub, f"{len(resultados)} archivos"
        hoja[f"B{fila}"].font = hoja[f"C{fila}"].font = Font(bold=True, italic=True)
        hoja[f"C{fila}"].fill = tot_r
        total += sub
        archivos += len(resultados)
        fila += 2
    hoja[f"B{fila}"], hoja[f"C{fila}"], hoja[f"D{fila}"] = "TOTAL GENERAL", total, f"{archivos} archivos"
    for c in ("B", "C", "D"):
        hoja[f"{c}{fila}"].font = tot_f
    hoja[f"C{fila}"].fill = tot_r
    for col, ancho in (("A", 22), ("B", 70), ("C", 20), ("D", 16)):
        hoja.column_dimensions[col].width = ancho
    info = libro.create_sheet("Información")
    info["A1"] = "Información del Reporte"
    info["A1"].font = Font(bold=True, size=14)
    for i, (k, v) in enumerate((("Fecha de generación:", dt.datetime.now().strftime("%d/%m/%Y %H:%M:%S")),
                                ("Tipo de búsqueda:", "Solo index.pdf" if solo_index else "Todos los PDFs"),
                                ("Total de grupos procesados:", len(filas)), ("Total de archivos:", archivos),
                                ("Total de páginas:", total), ("Generado por:", "Raphi · Inventario de páginas")), start=3):
        info[f"A{i}"], info[f"B{i}"] = k, v
    libro.save(destino)
    return total
