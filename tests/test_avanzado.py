"""R7: tachar de verdad, resaltar, notas, estampar, formularios, aplanar, comparar, firmas, tablas y comprobantes."""
from __future__ import annotations

import pytest

from raphi import motor
from raphi.motor.pdf_avanzado import leer_datos


def correr(op, entradas, salida, **params):
    inf = motor.ejecutar(op, entradas, params, destino=salida)
    assert inf.ok, [e.error for e in inf.elementos]
    return inf


@pytest.fixture
def con_datos(tmp_path):
    import pymupdf
    doc = pymupdf.open()
    p = doc.new_page()
    p.insert_text((72, 100), "Solicitante: Ana Quispe, DNI 45678912, correo ana@ejemplo.pe", fontsize=11)
    p.insert_text((72, 130), "Celular 987654321. Empresa RUC 20123456789. Texto que se queda.", fontsize=11)
    ruta = tmp_path / "datos.pdf"
    doc.save(ruta)
    return ruta


def texto(ruta):
    import pymupdf
    with pymupdf.open(ruta) as d:
        return "\n".join(p.get_text() for p in d)


def test_tachar_borra_de_verdad(con_datos, tmp_path):
    t = correr("tachar", [con_datos], tmp_path / "s", terminos="Ana Quispe", dni=True, ruc=True, correo=True, celular=True).archivos[0]
    queda = texto(t)
    for secreto in ("Ana Quispe", "45678912", "20123456789", "ana@ejemplo.pe", "987654321"):
        assert secreto not in queda
    assert "Texto que se queda" in queda
    inf = motor.ejecutar("tachar", [con_datos], {}, destino=tmp_path / "s")
    assert not inf.ok and "Diga qué tachar" in inf.elementos[0].error


def test_resaltar_nota_aplanar(con_datos, tmp_path):
    import pymupdf
    r = correr("resaltar", [con_datos], tmp_path / "s", terminos="Empresa").archivos[0]
    with pymupdf.open(r) as d:
        assert [a.type[1] for a in d[0].annots()] == ["Highlight"]
    n = correr("nota", [r], tmp_path / "s", texto="Revisar el plazo", autor="Mesa").archivos[0]
    a = correr("aplanar", [n], tmp_path / "s").archivos[0]
    with pymupdf.open(a) as d:
        assert not list(d[0].annots()) and "Revisar el plazo" in d[0].get_text()
    assert correr("resaltar", [con_datos], tmp_path / "s", terminos="inexistente").elementos[0].resultado.sin_cambio


def test_estampar_firma(corpus, tmp_path):
    import pymupdf
    e = correr("estampar_imagen", [corpus["texto"], corpus["png"]], tmp_path / "s", paginas="z").archivos[0]
    with pymupdf.open(e) as d:
        assert not d[0].get_images() and d[2].get_images()
    inf = motor.ejecutar("estampar_imagen", [corpus["texto"], corpus["png"], corpus["jpg"]], {}, destino=tmp_path / "s")
    assert not inf.ok and "UNA imagen" in inf.elementos[0].error


def test_formulario(tmp_path):
    import json
    import pymupdf
    doc = pymupdf.open()
    p = doc.new_page()
    for nombre, y in (("nombre", 100), ("cargo", 140)):
        w = pymupdf.Widget()
        w.field_name, w.field_type, w.rect = nombre, pymupdf.PDF_WIDGET_TYPE_TEXT, pymupdf.Rect(72, y, 300, y + 20)
        p.add_widget(w)
    ruta = tmp_path / "form.pdf"
    doc.save(ruta)
    campos = correr("formulario", [ruta], tmp_path / "s").elementos[0].resultado.datos["campos"]
    assert {c["nombre"] for c in campos} == {"nombre", "cargo"}
    f = correr("formulario", [ruta], tmp_path / "s", valores=json.dumps({"nombre": "Ana"}), aplanar=True).archivos[0]
    assert "Ana" in texto(f)
    inf = motor.ejecutar("formulario", [ruta], {"valores": '{"otro": 1}'}, destino=tmp_path / "s")
    assert not inf.ok and "otro" in inf.elementos[0].error


def test_comparar(corpus, tmp_path):
    import pymupdf
    nuevo = tmp_path / "v2.pdf"
    with pymupdf.open(corpus["texto"]) as d:
        d[1].insert_text((72, 300), "Artículo nuevo añadido", fontsize=12)
        d.save(nuevo)
    c = correr("comparar", [corpus["texto"], nuevo], tmp_path / "s").elementos[0].resultado
    assert c.datos["paginas_distintas"] == [2]
    informe = c.archivos[0].read_text()
    assert "<ins>Artículo nuevo añadido</ins>" in informe
    iguales = correr("comparar", [corpus["texto"], corpus["texto"]], tmp_path / "s").elementos[0].resultado
    assert iguales.datos["paginas_distintas"] == []


def test_verificar_firma_hecha_con_un_certificado_de_prueba(corpus, tmp_path):
    import datetime as dt
    from cryptography import x509
    from cryptography.hazmat.primitives import hashes, serialization
    from cryptography.hazmat.primitives.asymmetric import rsa
    from cryptography.x509.oid import NameOID
    from pyhanko.pdf_utils.incremental_writer import IncrementalPdfFileWriter
    from pyhanko.sign import signers
    from pyhanko.sign.fields import SigFieldSpec, append_signature_field

    clave = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    nombre = x509.Name([x509.NameAttribute(NameOID.COMMON_NAME, "Firmante de Prueba")])
    ahora = dt.datetime.now(dt.timezone.utc)
    cert = (x509.CertificateBuilder().subject_name(nombre).issuer_name(nombre).public_key(clave.public_key())
            .serial_number(1).not_valid_before(ahora - dt.timedelta(days=1)).not_valid_after(ahora + dt.timedelta(days=30))
            .sign(clave, hashes.SHA256()))
    (tmp_path / "k.pem").write_bytes(clave.private_bytes(serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8,
                                                         serialization.NoEncryption()))
    (tmp_path / "c.pem").write_bytes(cert.public_bytes(serialization.Encoding.PEM))
    firmante = signers.SimpleSigner.load(str(tmp_path / "k.pem"), str(tmp_path / "c.pem"))
    firmado = tmp_path / "firmado.pdf"
    with open(corpus["texto"], "rb") as f:
        w = IncrementalPdfFileWriter(f)
        append_signature_field(w, SigFieldSpec("Firma1"))
        with open(firmado, "wb") as salida:
            signers.sign_pdf(w, signers.PdfSignatureMetadata(field_name="Firma1"), signer=firmante, output=salida)
    f = correr("firmas", [firmado], tmp_path / "s").elementos[0].resultado.datos["firmas"]
    assert len(f) == 1 and f[0]["integra"] and "Firmante de Prueba" in f[0]["firmante"] and "sin raíces" in f[0]["cadena"]
    # Alterado después de firmar: deja de estar íntegro.
    alterado = tmp_path / "alterado.pdf"
    datos = bytearray(firmado.read_bytes())
    i = datos.find(b"Lorem")
    datos[i:i + 5] = b"LOREM"
    alterado.write_bytes(bytes(datos))
    g = motor.ejecutar("firmas", [alterado], {}, destino=tmp_path / "s").elementos[0]
    assert not g.ok or not g.resultado.datos["firmas"][0]["integra"]


def test_tablas_a_excel(tmp_path):
    import pymupdf
    doc = pymupdf.open()
    p = doc.new_page()
    for i in range(4):
        for j in range(3):
            r = pymupdf.Rect(72 + j * 100, 100 + i * 24, 172 + j * 100, 124 + i * 24)
            p.draw_rect(r, color=(0, 0, 0), width=0.8)
            p.insert_text((r.x0 + 4, r.y0 + 16), f"c{i}{j}", fontsize=10)
    ruta = tmp_path / "tabla.pdf"
    doc.save(ruta)
    x = correr("tablas_a_excel", [ruta], tmp_path / "s").archivos[0]
    from openpyxl import load_workbook
    hoja = load_workbook(x).worksheets[0]
    assert hoja.max_row == 4 and hoja["C4"].value == "c32"


def test_leer_comprobante():
    d = leer_datos("BOLETA DE VENTA ELECTRÓNICA\nRUC 20601234567\nB001 - 00004521\nFecha de emisión: 15/10/2026\n"
                   "OP. GRAVADA S/ 84.75\nIGV 15.25\nIMPORTE TOTAL: S/ 100,00")
    assert d == {"ruc": "20601234567", "serie": "B001", "numero": 4521, "fecha": "2026-10-15", "total": 100.0}
    assert leer_datos("Total 1.234,50")["total"] == 1234.5
